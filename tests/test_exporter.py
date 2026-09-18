import io
from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app import db, exporter, indexer
from tests.conftest import add_indexed_file


def _make_pdf(path: Path, pages=2):
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page(width=720, height=405).insert_text((50, 100), f"page {i}")
    doc.save(str(path))
    doc.close()


def _make_pptx(path: Path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "native text"
    for _ in range(2):  # two chart slides -> two image fallbacks from the same deck
        c = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("s", (1, 2))
        c.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
    prs.save(str(path))


@pytest.fixture()
def library(tmp_db, tmp_path, monkeypatch):
    pptx_path, pdf_path = tmp_path / "deck.pptx", tmp_path / "doc.pdf"
    _make_pptx(pptx_path)
    _make_pdf(pdf_path, pages=3)
    sid = db.add_source(str(tmp_path), "Dom")
    ids = {
        "pptx": add_indexed_file(sid, pptx_path, [("a", "x")] * 3, thumbs=False),
        "pdf": add_indexed_file(sid, pdf_path, [("a", "x")] * 3, ext="pdf", thumbs=False),
    }
    # soffice stand-in: a real PDF with one page per source slide
    calls = []

    def fake_convert(path, out_dir, expected_pages=None):
        calls.append(str(path))
        out = out_dir / (Path(path).stem + ".pdf")
        _make_pdf(out, pages=3)
        return out

    monkeypatch.setattr(indexer, "_convert_to_pdf", fake_convert)
    return ids, calls


def _reload(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def test_native_slide_stays_editable_and_dividers_are_added(library):
    ids, _ = library
    data = exporter.build_deck([{"name": "Intro", "slides": [{"file_id": ids["pptx"], "slide_index": 0}]}])
    prs = _reload(data)
    assert len(prs.slides) == 2  # divider + slide
    divider_text = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert "Intro" in divider_text
    assert "native text" in [s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame]


def test_no_dividers_option(library):
    ids, _ = library
    data = exporter.build_deck([{"name": "X", "slides": [{"file_id": ids["pptx"], "slide_index": 0}]}], add_dividers=False)
    assert len(_reload(data).slides) == 1


def test_pdf_and_chart_slides_become_full_bleed_pictures(library):
    ids, _ = library
    refs = [{"file_id": ids["pdf"], "slide_index": 1}, {"file_id": ids["pptx"], "slide_index": 1}]
    prs = _reload(exporter.build_deck([{"name": "C", "slides": refs}], add_dividers=False))
    for slide in prs.slides:
        shapes = list(slide.shapes)
        assert len(shapes) == 1 and shapes[0].shape_type == 13
        assert shapes[0].width == exporter.SLIDE_WIDTH


def test_each_source_deck_is_converted_once_per_export(library):
    ids, calls = library
    refs = [{"file_id": ids["pptx"], "slide_index": 1}, {"file_id": ids["pptx"], "slide_index": 2}]
    exporter.build_deck([{"name": "C", "slides": refs}], add_dividers=False)
    assert len(calls) == 1


def test_missing_file_row_is_skipped(library):
    data = exporter.build_deck([{"name": "C", "slides": [{"file_id": 9999, "slide_index": 0}]}], add_dividers=False)
    assert len(_reload(data).slides) == 0
