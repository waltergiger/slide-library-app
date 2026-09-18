import subprocess
import sys
import textwrap
from pathlib import Path, PureWindowsPath

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app import db, folder_picker, indexer, main, soffice

REPO = Path(__file__).resolve().parent.parent


# ---- data dir ---------------------------------------------------------------

def test_data_dir_can_be_relocated_with_env(tmp_path):
    target = tmp_path / "profile" / "SlideLibrary" / "data"
    out = subprocess.run(
        [sys.executable, "-c", "from app import db; print(db.DATA_DIR); print(db.THUMB_DIR.exists())"],
        cwd=REPO, env={"SLIDELIB_DATA_DIR": str(target), "PATH": ""}, capture_output=True, text=True, check=True,
    ).stdout.split()
    assert out == [str(target), "True"]  # nested dirs are created, thumbnails dir included


# ---- soffice discovery --------------------------------------------------------

def test_windows_candidates_cover_standard_install_locations():
    env = {"ProgramFiles": r"C:\Program Files", "ProgramFiles(x86)": r"C:\Program Files (x86)",
           "LOCALAPPDATA": r"C:\Users\me\AppData\Local"}
    got = [PureWindowsPath(p) for p in soffice.candidate_paths("win32", env)]
    assert PureWindowsPath(r"C:\Program Files\LibreOffice\program\soffice.exe") in got
    assert PureWindowsPath(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe") in got
    assert any("AppData" in str(p) for p in got)
    assert soffice.candidate_paths("win32", {}) == []  # no env vars -> no bogus relative paths


def test_macos_and_linux_candidates():
    assert soffice.candidate_paths("darwin", {}) == ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    assert "/usr/bin/soffice" in soffice.candidate_paths("linux", {})


def test_find_soffice_order_env_override_then_path_then_known_locations(tmp_path, monkeypatch):
    fake = tmp_path / "soffice"
    fake.write_text("")
    monkeypatch.setenv("SLIDELIB_SOFFICE", str(fake))
    assert soffice.find_soffice() == str(fake)

    monkeypatch.setenv("SLIDELIB_SOFFICE", str(tmp_path / "nope"))
    assert soffice.find_soffice() is None  # an explicit override that doesn't exist is an error, not a silent fallback

    monkeypatch.delenv("SLIDELIB_SOFFICE")
    monkeypatch.setattr(soffice.shutil, "which", lambda name: "/opt/x/soffice")
    assert soffice.find_soffice() == "/opt/x/soffice"

    monkeypatch.setattr(soffice.shutil, "which", lambda name: None)
    monkeypatch.setattr(soffice, "candidate_paths", lambda platform, env: [str(fake)])
    assert soffice.find_soffice() == str(fake)
    monkeypatch.setattr(soffice, "candidate_paths", lambda platform, env: [])
    assert soffice.find_soffice() is None


# ---- indexing without LibreOffice ---------------------------------------------

def _pptx(path: Path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "quarterly wealth review"
    prs.save(str(path))


def test_pptx_is_still_indexed_and_searchable_when_libreoffice_is_missing(tmp_db, tmp_path, monkeypatch):
    lib = tmp_path / "lib"
    lib.mkdir()
    _pptx(lib / "Deck.pptx")
    monkeypatch.setattr(soffice, "find_soffice", lambda: None)
    monkeypatch.setattr(indexer, "_warned_no_soffice", False)

    sid = db.add_source(str(lib), "Dom")
    indexer.index_source(sid)

    assert db.get_source(sid)["status"] == "indexed"
    hits = db.search_slides("wealth")
    assert len(hits) == 1 and hits[0]["thumb_file"] is None  # text indexed, thumbnail simply absent


def test_windows_profile_url_is_a_valid_file_uri():
    assert PureWindowsPath(r"C:\Users\Jane Doe\AppData\Local\Temp\lo_profile").as_uri() == \
        "file:///C:/Users/Jane%20Doe/AppData/Local/Temp/lo_profile"


def test_convert_passes_valid_uri_and_flags_to_soffice(tmp_path, monkeypatch):
    seen = {}
    monkeypatch.setattr(soffice, "find_soffice", lambda: "/x/soffice")
    monkeypatch.setattr(indexer.subprocess, "run", lambda cmd, **kw: seen.update(cmd=cmd, kw=kw) or
                        subprocess.CompletedProcess(cmd, 1, b"", b"boom"))
    assert indexer._convert_to_pdf(tmp_path / "a.pptx", tmp_path) is None
    assert seen["cmd"][0] == "/x/soffice"
    assert any(a.startswith("-env:UserInstallation=file:///") for a in seen["cmd"])


# ---- health endpoint ------------------------------------------------------------

def test_health_reports_libreoffice_presence(tmp_db, monkeypatch):
    with TestClient(main.app, base_url="http://127.0.0.1:8420") as c:
        monkeypatch.setattr(soffice, "find_soffice", lambda: "/x/soffice")
        assert c.get("/api/health").json() == {"libreoffice": True, "libreoffice_path": "/x/soffice"}
        monkeypatch.setattr(soffice, "find_soffice", lambda: None)
        assert c.get("/api/health").json() == {"libreoffice": False, "libreoffice_path": None}


# ---- Windows folder picker -------------------------------------------------------

def _proc(code=0, out="", err=""):
    return subprocess.CompletedProcess([], code, out, err)


def test_windows_picker_uses_powershell_sta_and_returns_path(monkeypatch):
    seen = {}
    monkeypatch.setattr(folder_picker.sys, "platform", "win32")
    monkeypatch.setattr(folder_picker.subprocess, "run", lambda cmd, **kw: seen.update(cmd=cmd, kw=kw) or _proc(out="\ufeffD:\\Decks\\Strategy\r\n"))
    assert folder_picker.pick_folder() == "D:\\Decks\\Strategy"   # BOM and CRLF stripped
    assert seen["cmd"][0] == "powershell" and "-STA" in seen["cmd"]  # STA is required for WinForms dialogs
    assert seen["cmd"][0] != sys.executable                          # never re-launch the (frozen) app itself
    assert seen["kw"]["encoding"] == "utf-8"
    assert "FolderBrowserDialog" in seen["cmd"][-1]


def test_windows_picker_cancel_and_failure(monkeypatch):
    monkeypatch.setattr(folder_picker.sys, "platform", "win32")
    monkeypatch.setattr(folder_picker.subprocess, "run", lambda cmd, **kw: _proc(out=""))
    assert folder_picker.pick_folder() is None
    monkeypatch.setattr(folder_picker.subprocess, "run", lambda cmd, **kw: _proc(1, err="Add-Type failed"))
    with pytest.raises(folder_picker.PickerUnavailable, match="Add-Type"):
        folder_picker.pick_folder()


def test_windows_prompt_is_escaped_for_powershell(monkeypatch):
    seen = {}
    monkeypatch.setattr(folder_picker.sys, "platform", "win32")
    monkeypatch.setattr(folder_picker.subprocess, "run", lambda cmd, **kw: seen.update(cmd=cmd) or _proc())
    folder_picker.pick_folder("Bob's decks")
    assert "'Bob''s decks'" in seen["cmd"][-1]


def test_frozen_app_never_spawns_itself_as_a_tk_helper(monkeypatch):
    monkeypatch.setattr(folder_picker.sys, "platform", "linux")
    monkeypatch.setattr(folder_picker.sys, "frozen", True, raising=False)
    monkeypatch.setattr(folder_picker.subprocess, "run", lambda *a, **k: pytest.fail("must not spawn"))
    with pytest.raises(folder_picker.PickerUnavailable):
        folder_picker.pick_folder()
