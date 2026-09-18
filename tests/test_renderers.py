import subprocess
import sys
import types
from pathlib import Path, PureWindowsPath

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app import db, indexer, main, renderers
from app.renderers import libreoffice, manager, powerpoint
from app.renderers.base import Engine, RenderError


def _make_pdf(path: Path, pages: int):
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page(width=720, height=405).insert_text((50, 100), f"page {i}")
    doc.save(str(path))
    doc.close()
    return path


def _pages(pdf: Path) -> int:
    doc = pymupdf.open(str(pdf))
    try:
        return doc.page_count
    finally:
        doc.close()


def _make_pptx(path: Path, slides=3, hidden=()):
    prs = Presentation()
    for i in range(slides):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = f"slide {i + 1} content"
        if i in hidden:
            s._element.set("show", "0")
    prs.save(str(path))
    return path


@pytest.fixture(autouse=True)
def clean_state(monkeypatch):
    monkeypatch.delenv("SLIDELIB_ENGINE", raising=False)
    manager.reset()
    yield
    manager.reset()


class Fake(Engine):
    def __init__(self, id_, installed=True, fail=None, pages=3):
        self.id, self.label = id_, id_.title()
        self._installed, self._fail, self.pages, self.calls = installed, fail, pages, 0

    def installed(self):
        return self._installed

    def convert(self, path, out_dir):
        self.calls += 1
        if self._fail:
            raise self._fail
        return _make_pdf(out_dir / (path.stem + ".pdf"), self.pages)


def use(monkeypatch, *engines):
    monkeypatch.setattr(manager, "_engines", lambda: list(engines))
    return engines


# ---- selection & fallback --------------------------------------------------------

def test_powerpoint_is_tried_first_and_libreoffice_is_not_touched_when_it_works(monkeypatch, tmp_path):
    pp, lo = use(monkeypatch, Fake("powerpoint"), Fake("libreoffice"))
    assert manager.convert_to_pdf(tmp_path / "a.pptx", tmp_path).name == "a.pdf"
    assert (pp.calls, lo.calls) == (1, 0)
    assert manager.status()["active"] == "powerpoint"


def test_libreoffice_takes_over_when_powerpoint_is_not_installed(monkeypatch, tmp_path):
    pp, lo = use(monkeypatch, Fake("powerpoint", installed=False), Fake("libreoffice"))
    assert manager.convert_to_pdf(tmp_path / "a.pptx", tmp_path) is not None
    assert (pp.calls, lo.calls) == (0, 1)
    st = manager.status()
    assert st["active"] == "libreoffice" and st["engines"][0]["installed"] is False


def test_a_failed_file_falls_through_to_the_next_engine(monkeypatch, tmp_path):
    pp, lo = use(monkeypatch, Fake("powerpoint", fail=RenderError("password protected")), Fake("libreoffice"))
    assert manager.convert_to_pdf(tmp_path / "a.pptx", tmp_path) is not None
    assert (pp.calls, lo.calls) == (1, 1)
    assert manager.status()["active"] == "powerpoint"  # one bad file doesn't switch the engine off


def test_systemic_failure_switches_the_engine_off_immediately_and_reports_why(monkeypatch, tmp_path):
    pp, lo = use(monkeypatch, Fake("powerpoint", fail=RenderError("macOS blocked automation", systemic=True)), Fake("libreoffice"))
    for name in ("a", "b", "c"):
        assert manager.convert_to_pdf(tmp_path / f"{name}.pptx", tmp_path) is not None
    assert pp.calls == 1 and lo.calls == 3  # PowerPoint is not retried (and waited on) for every file
    st = manager.status()
    assert st["active"] == "libreoffice"
    assert st["engines"][0]["problem"] == "macOS blocked automation" and st["engines"][0]["usable"] is False


def test_repeated_ordinary_failures_also_switch_the_engine_off(monkeypatch, tmp_path):
    pp, lo = use(monkeypatch, Fake("powerpoint", fail=RenderError("boom")), Fake("libreoffice"))
    for i in range(manager.MAX_CONSECUTIVE_FAILURES + 2):
        manager.convert_to_pdf(tmp_path / f"{i}.pptx", tmp_path)
    assert pp.calls == manager.MAX_CONSECUTIVE_FAILURES
    manager.reset()
    manager.convert_to_pdf(tmp_path / "again.pptx", tmp_path)
    assert pp.calls == manager.MAX_CONSECUTIVE_FAILURES + 1  # reset() re-enables it


def test_a_success_resets_the_failure_streak(monkeypatch, tmp_path):
    flaky = Fake("powerpoint", fail=RenderError("boom"))
    use(monkeypatch, flaky, Fake("libreoffice"))
    for _ in range(2):
        manager.convert_to_pdf(tmp_path / "x.pptx", tmp_path)
    flaky._fail = None
    manager.convert_to_pdf(tmp_path / "ok.pptx", tmp_path)
    flaky._fail = RenderError("boom")
    for _ in range(2):
        manager.convert_to_pdf(tmp_path / "y.pptx", tmp_path)
    assert manager.status()["engines"][0]["usable"] is True  # 2 + 2 failures, but never 3 in a row


def test_no_engine_installed_returns_none_and_status_says_so(monkeypatch, tmp_path):
    use(monkeypatch, Fake("powerpoint", installed=False), Fake("libreoffice", installed=False))
    assert manager.convert_to_pdf(tmp_path / "a.pptx", tmp_path) is None
    assert not manager.has_usable_engine()
    st = manager.status()
    assert st["active"] is None and [e["installed"] for e in st["engines"]] == [False, False]


def test_installed_but_switched_off_engine_means_no_active_renderer(monkeypatch, tmp_path):
    use(monkeypatch, Fake("powerpoint", fail=RenderError("blocked", systemic=True)), Fake("libreoffice", installed=False))
    manager.convert_to_pdf(tmp_path / "a.pptx", tmp_path)
    st = manager.status()
    assert st["active"] is None
    assert st["engines"][0]["installed"] is True and st["engines"][0]["problem"] == "blocked"


def test_env_can_force_a_single_engine(monkeypatch):
    monkeypatch.setenv("SLIDELIB_ENGINE", "libreoffice")
    assert [e.id for e in manager._engines()] == ["libreoffice"]
    monkeypatch.setenv("SLIDELIB_ENGINE", "powerpoint")
    assert [e.id for e in manager._engines()] in (["powerpoint"], [])  # [] on Linux: no PowerPoint there
    monkeypatch.setenv("SLIDELIB_ENGINE", "auto")
    assert [e.id for e in manager._engines()][-1] == "libreoffice"


def test_platform_engine_lineup(monkeypatch):
    monkeypatch.setattr(sys, "platform", "darwin")
    assert isinstance(powerpoint.powerpoint_engine(), powerpoint.PowerPointMacEngine)
    monkeypatch.setattr(sys, "platform", "win32")
    assert isinstance(powerpoint.powerpoint_engine(), powerpoint.PowerPointWindowsEngine)
    monkeypatch.setattr(sys, "platform", "linux")
    assert powerpoint.powerpoint_engine() is None


# ---- LibreOffice engine ------------------------------------------------------------

def test_soffice_candidates_and_lookup_order(tmp_path, monkeypatch):
    env = {"ProgramFiles": r"C:\Program Files", "ProgramFiles(x86)": r"C:\Program Files (x86)", "LOCALAPPDATA": r"C:\Users\me\AppData\Local"}
    got = [PureWindowsPath(p) for p in libreoffice.candidate_paths("win32", env)]
    assert PureWindowsPath(r"C:\Program Files\LibreOffice\program\soffice.exe") in got and any("AppData" in str(p) for p in got)
    assert libreoffice.candidate_paths("win32", {}) == []
    assert "/Applications/LibreOffice.app/Contents/MacOS/soffice" in libreoffice.candidate_paths("darwin", {})

    fake = tmp_path / "soffice"
    fake.write_text("")
    monkeypatch.setenv("SLIDELIB_SOFFICE", str(fake))
    assert libreoffice.find_soffice() == str(fake)
    monkeypatch.setenv("SLIDELIB_SOFFICE", str(tmp_path / "missing"))
    assert libreoffice.find_soffice() is None  # an explicit override that is wrong is an error, not a silent fallback
    monkeypatch.delenv("SLIDELIB_SOFFICE")
    monkeypatch.setattr(libreoffice.shutil, "which", lambda n: None)
    monkeypatch.setattr(libreoffice, "candidate_paths", lambda p, e: [str(fake)])
    assert libreoffice.find_soffice() == str(fake)


def test_libreoffice_command_keeps_hidden_slides_and_uses_a_valid_profile_uri(tmp_path, monkeypatch):
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: "/x/soffice")
    calls = []

    def fake_run(cmd, **kw):
        calls.append((cmd, kw))
        (tmp_path / "deck.pdf").write_bytes(b"%PDF")
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(libreoffice.subprocess, "run", fake_run)
    out = libreoffice.LibreOfficeEngine().convert(tmp_path / "deck.pptx", tmp_path)
    assert out == tmp_path / "deck.pdf" and len(calls) == 1
    cmd = calls[0][0]
    assert cmd[0] == "/x/soffice" and "ExportHiddenSlides" in cmd[cmd.index("--convert-to") + 1]
    assert any(a.startswith("-env:UserInstallation=file:///") for a in cmd)
    assert PureWindowsPath(r"C:\Users\Jane Doe\Temp\lo_profile").as_uri() == "file:///C:/Users/Jane%20Doe/Temp/lo_profile"


def test_libreoffice_retries_with_the_plain_filter_when_the_option_is_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: "/x/soffice")
    filters = []

    def fake_run(cmd, **kw):
        f = cmd[cmd.index("--convert-to") + 1]
        filters.append(f)
        if f == "pdf":  # old LibreOffice: only the plain filter works
            (tmp_path / "deck.pdf").write_bytes(b"%PDF")
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(libreoffice.subprocess, "run", fake_run)
    assert libreoffice.LibreOfficeEngine().convert(tmp_path / "deck.pptx", tmp_path).exists()
    assert len(filters) == 2 and filters[1] == "pdf"


def test_libreoffice_errors_are_classified(tmp_path, monkeypatch):
    eng = libreoffice.LibreOfficeEngine()
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: None)
    with pytest.raises(RenderError) as e:
        eng.convert(tmp_path / "a.pptx", tmp_path)
    assert e.value.systemic

    monkeypatch.setattr(libreoffice, "find_soffice", lambda: "/x/soffice")
    monkeypatch.setattr(libreoffice.subprocess, "run", lambda *a, **k: subprocess.CompletedProcess(a, 0))
    with pytest.raises(RenderError, match="no PDF"):
        eng.convert(tmp_path / "a.pptx", tmp_path)

    def hang(*a, **k):
        raise subprocess.TimeoutExpired("soffice", 1)
    monkeypatch.setattr(libreoffice.subprocess, "run", hang)
    with pytest.raises(RenderError, match="timed out"):
        eng.convert(tmp_path / "a.pptx", tmp_path)


@pytest.mark.skipif(libreoffice.find_soffice() is None, reason="LibreOffice not installed")
def test_real_libreoffice_renders_hidden_slides_so_page_n_is_slide_n(tmp_path):
    deck = _make_pptx(tmp_path / "hidden.pptx", slides=4, hidden=(1,))
    pdf = libreoffice.LibreOfficeEngine().convert(deck, tmp_path)
    doc = pymupdf.open(str(pdf))
    assert [p.get_text().strip() for p in doc] == [f"slide {i} content" for i in (1, 2, 3, 4)]


# ---- PowerPoint on macOS -------------------------------------------------------------

@pytest.fixture()
def mac(monkeypatch, tmp_path):
    container = tmp_path / "Containers" / "com.microsoft.Powerpoint" / "Data"
    container.mkdir(parents=True)
    monkeypatch.setattr(powerpoint, "_MAC_CONTAINER", container)
    return container


def test_mac_detects_the_app_in_standard_locations(tmp_path, monkeypatch):
    apps = tmp_path / "Applications"
    monkeypatch.setattr(powerpoint, "_MAC_APP_DIRS", (apps,))
    monkeypatch.setattr(powerpoint, "_mdfind_cache", (10**12, None))  # fresh cache: Spotlight fallback not consulted
    assert powerpoint.mac_powerpoint_path() is None
    (apps / "Microsoft PowerPoint.app").mkdir(parents=True)
    assert powerpoint.mac_powerpoint_path() == str(apps / "Microsoft PowerPoint.app")


def test_mac_converts_a_copy_inside_the_sandbox_container_and_leaves_the_original_alone(mac, tmp_path, monkeypatch):
    (tmp_path / "src").mkdir()
    original = _make_pptx(tmp_path / "src" / "Q3 Board – Kickoff (v2).pptx")
    before = original.read_bytes()
    seen = {}

    def fake_osascript(cmd, **kw):
        src, out = Path(cmd[-2]), Path(cmd[-1])
        seen.update(cmd=cmd, src=src, out=out, src_existed=src.is_file())
        assert mac in src.parents and mac in out.parents        # sandbox: only its own container is reachable
        assert src.name == "deck.pptx"                          # plain name: nothing to escape in AppleScript
        _make_pdf(out, 3)
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(powerpoint.subprocess, "run", fake_osascript)
    dest_dir = tmp_path / "out"
    dest_dir.mkdir()
    pdf = powerpoint.PowerPointMacEngine().convert(original, dest_dir)

    assert pdf == dest_dir / "Q3 Board – Kickoff (v2).pdf" and pdf.is_file()
    assert seen["cmd"][0] == "osascript" and seen["src_existed"]
    script = "\n".join(seen["cmd"][i + 1] for i, a in enumerate(seen["cmd"]) if a == "-e")
    assert "save pres in POSIX file outPath as save as PDF" in script
    assert "wasRunning" in script and "quit" in script           # only quits an instance it started itself
    assert original.read_bytes() == before                        # original untouched
    assert not (mac / "SlideLibrary" / seen["src"].parent.name).exists()  # work dir cleaned up


def test_mac_permission_denied_is_systemic_with_actionable_text(mac, tmp_path, monkeypatch):
    src = _make_pptx(tmp_path / "a.pptx")
    monkeypatch.setattr(powerpoint.subprocess, "run", lambda cmd, **k: subprocess.CompletedProcess(
        cmd, 1, "", "execution error: Not authorized to send Apple events to Microsoft PowerPoint. (-1743)"))
    with pytest.raises(RenderError) as e:
        powerpoint.PowerPointMacEngine().convert(src, tmp_path)
    assert e.value.systemic and "Automation" in str(e.value)


def test_mac_other_failures_are_per_file_and_timeouts_point_at_dialogs(mac, tmp_path, monkeypatch):
    src = _make_pptx(tmp_path / "a.pptx")
    eng = powerpoint.PowerPointMacEngine()
    monkeypatch.setattr(powerpoint.subprocess, "run", lambda cmd, **k: subprocess.CompletedProcess(cmd, 1, "", "file is corrupt"))
    with pytest.raises(RenderError, match="corrupt") as e:
        eng.convert(src, tmp_path)
    assert not e.value.systemic

    def hang(cmd, **k):
        raise subprocess.TimeoutExpired("osascript", 1)
    monkeypatch.setattr(powerpoint.subprocess, "run", hang)
    with pytest.raises(RenderError, match="dialog"):
        eng.convert(src, tmp_path)

    monkeypatch.setattr(powerpoint.subprocess, "run", lambda cmd, **k: subprocess.CompletedProcess(cmd, 0, "", ""))
    with pytest.raises(RenderError, match="no PDF"):
        eng.convert(src, tmp_path)
    assert list((mac / "SlideLibrary").iterdir()) == []           # nothing left behind in any failure case


# ---- PowerPoint on Windows -----------------------------------------------------------

def test_windows_uses_com_via_powershell_with_paths_in_env_vars(tmp_path, monkeypatch):
    src = _make_pptx(tmp_path / "My deck (final) 'v2'.pptx")
    seen = {}

    def fake_ps(cmd, **kw):
        seen.update(cmd=cmd, env=kw["env"])
        _make_pdf(Path(kw["env"]["SLIDELIB_OUT"]), 3)
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(powerpoint.subprocess, "run", fake_ps)
    pdf = powerpoint.PowerPointWindowsEngine().convert(src, tmp_path)
    assert pdf == tmp_path / "My deck (final) 'v2'.pdf"
    assert seen["cmd"][0] == "powershell" and "-STA" in seen["cmd"]
    assert seen["env"]["SLIDELIB_IN"] == str(src)                 # via env: quotes/parentheses in names can't break the script
    script = seen["cmd"][-1]
    assert "Presentations.Open($env:SLIDELIB_IN, -1, -1, 0)" in script       # read-only, windowless
    assert "ExportAsFixedFormat($env:SLIDELIB_OUT, 2, 2, 0, 1, 1, -1)" in script   # PDF, print quality, hidden slides included
    assert "$wasRunning" in script and "-not $wasRunning" in script          # never quits the user's own PowerPoint


def test_windows_com_not_registered_is_systemic_and_other_errors_are_per_file(tmp_path, monkeypatch):
    src = _make_pptx(tmp_path / "a.pptx")
    eng = powerpoint.PowerPointWindowsEngine()
    monkeypatch.setattr(powerpoint.subprocess, "run", lambda cmd, **k: subprocess.CompletedProcess(
        cmd, 1, "", "Retrieving the COM class factory ... 80040154 Class not registered"))
    with pytest.raises(RenderError) as e:
        eng.convert(src, tmp_path)
    assert e.value.systemic and "Store" in str(e.value)
    monkeypatch.setattr(powerpoint.subprocess, "run", lambda cmd, **k: subprocess.CompletedProcess(cmd, 1, "", "The file is password protected"))
    with pytest.raises(RenderError, match="password") as e:
        eng.convert(src, tmp_path)
    assert not e.value.systemic


def test_windows_install_detection_via_registry(monkeypatch):
    class Key:
        def __enter__(self): return self
        def __exit__(self, *a): return False

    reg = types.SimpleNamespace(HKEY_CLASSES_ROOT=0, OpenKey=lambda root, path: Key() if path == r"PowerPoint.Application\CLSID" else (_ for _ in ()).throw(OSError()))
    monkeypatch.setitem(sys.modules, "winreg", reg)
    assert powerpoint.windows_powerpoint_installed() is True
    monkeypatch.setitem(sys.modules, "winreg", types.SimpleNamespace(HKEY_CLASSES_ROOT=0, OpenKey=lambda r, p: (_ for _ in ()).throw(OSError())))
    assert powerpoint.windows_powerpoint_installed() is False
    monkeypatch.setitem(sys.modules, "winreg", None)   # not Windows: import fails
    assert powerpoint.windows_powerpoint_installed() is False


# ---- indexer integration ---------------------------------------------------------------

def _library(tmp_path, **kw):
    lib = tmp_path / "lib"
    lib.mkdir()
    _make_pptx(lib / "Deck.pptx", **kw)
    return lib


def test_deck_is_indexed_text_only_without_a_renderer_then_gets_thumbnails_after_install(tmp_db, tmp_path, monkeypatch):
    lib = _library(tmp_path)
    use(monkeypatch, Fake("powerpoint", installed=False), Fake("libreoffice", installed=False))
    sid = db.add_source(str(lib), "Dom")
    indexer.index_source(sid)
    assert db.get_source(sid)["status"] == "indexed"
    assert [r["thumb_file"] for r in db.list_slides(db.list_decks(None, None)[0]["id"])] == [None] * 3
    assert db.search_slides("content")                             # search works without any renderer

    indexer.index_source(sid)                                      # still nothing to render with: left alone
    assert all(r["thumb_file"] is None for r in db.list_slides(db.list_decks(None, None)[0]["id"]))

    use(monkeypatch, Fake("powerpoint"))                           # user installs PowerPoint...
    indexer.index_source(sid)                                      # ...and re-indexes: the unchanged deck is rendered now
    thumbs = [r["thumb_file"] for r in db.list_slides(db.list_decks(None, None)[0]["id"])]
    assert all(thumbs) and all((db.THUMB_DIR / t).exists() for t in thumbs)


def test_favorites_survive_the_late_thumbnail_pass(tmp_db, tmp_path, monkeypatch):
    lib = _library(tmp_path)
    use(monkeypatch, Fake("libreoffice", installed=False))
    sid = db.add_source(str(lib), "Dom")
    indexer.index_source(sid)
    deck = db.list_decks(None, None)[0]
    db.set_slide_favorite(deck["id"], 1, True)
    use(monkeypatch, Fake("libreoffice"))
    indexer.index_source(sid)
    assert [bool(r["favorite"]) for r in db.list_slides(deck["id"])] == [False, True, False]


def test_a_pdf_with_the_wrong_page_count_yields_no_thumbnails_instead_of_wrong_ones(tmp_db, tmp_path, monkeypatch):
    lib = _library(tmp_path, slides=4, hidden=(1,))
    use(monkeypatch, Fake("libreoffice", pages=3))                 # engine that dropped the hidden slide
    sid = db.add_source(str(lib), "Dom")
    indexer.index_source(sid)
    assert all(r["thumb_file"] is None for r in db.list_slides(db.list_decks(None, None)[0]["id"]))


def test_export_image_fallback_refuses_a_mismatched_pdf(tmp_path, monkeypatch):
    deck = _make_pptx(tmp_path / "d.pptx", slides=4)
    use(monkeypatch, Fake("libreoffice", pages=3))                 # only engine drops a slide
    with indexer.PdfRenderCache() as cache, pytest.raises(RuntimeError, match="none produced a PDF matching"):
        cache.pdf_for(deck)


def test_engine_with_the_wrong_page_count_is_skipped_for_the_next_one(monkeypatch, tmp_path):
    pp, lo = use(monkeypatch, Fake("powerpoint", pages=3), Fake("libreoffice", pages=4))
    pdf = manager.convert_to_pdf(tmp_path / "hid.pptx", tmp_path, expected_pages=4)
    assert (pp.calls, lo.calls) == (1, 1) and _pages(pdf) == 4
    assert manager.convert_to_pdf(tmp_path / "x.pptx", tmp_path) is not None   # no expectation -> first engine accepted as-is


def test_export_error_names_the_missing_renderers(tmp_path, monkeypatch):
    deck = _make_pptx(tmp_path / "d.pptx")
    monkeypatch.setattr(indexer, "_convert_to_pdf", lambda p, o, n=None: None)
    with indexer.PdfRenderCache() as cache, pytest.raises(RuntimeError, match="PowerPoint or LibreOffice"):
        cache.pdf_for(deck)


# ---- API -------------------------------------------------------------------------------

def test_renderer_status_and_recheck_endpoints(tmp_db, monkeypatch, tmp_path):
    use(monkeypatch, Fake("powerpoint", fail=RenderError("blocked", systemic=True)), Fake("libreoffice", installed=False))
    manager.convert_to_pdf(tmp_path / "a.pptx", tmp_path)          # switches PowerPoint off
    with TestClient(main.app, base_url="http://127.0.0.1:8420") as c:
        st = c.get("/api/renderers").json()
        assert st["active"] is None and st["engines"][0]["problem"] == "blocked"
        assert set(st) == {"platform", "active", "engines"}
        st = c.post("/api/renderers/recheck").json()               # user fixed the permission and clicked "Check again"
        assert st["active"] == "powerpoint" and st["engines"][0]["problem"] is None
        assert c.post("/api/renderers/recheck", headers={"origin": "https://evil.example.com"}).status_code == 403
