import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.pptx_copy import _pick_blank_layout, copy_slide_native, slide_has_unsupported_content
from tests.conftest import png_bytes


def _source_with_picture_table_link():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Hello library"
    box.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/x"
    slide.shapes.add_picture(io.BytesIO(png_bytes("blue")), Inches(1), Inches(3))
    slide.shapes.add_table(2, 2, Inches(5), Inches(1), Inches(3), Inches(1)).table.cell(0, 0).text = "cell"
    return prs, slide


def _roundtrip(target):
    buf = io.BytesIO()
    target.save(buf)
    return Presentation(io.BytesIO(buf.getvalue()))


def test_copy_carries_text_picture_table_and_hyperlink():
    src, slide = _source_with_picture_table_link()
    target = Presentation()
    copy_slide_native(src, slide, target, _pick_blank_layout(target))

    reopened = _roundtrip(target)
    new = reopened.slides[0]
    texts = [s.text_frame.text for s in new.shapes if s.has_text_frame]
    assert "Hello library" in texts
    assert any(s.shape_type == 13 for s in new.shapes)  # PICTURE
    assert any(s.has_table for s in new.shapes)
    rels = {r.target_ref for r in new.part.rels.values() if r.is_external}
    assert "https://example.com/x" in rels
    pic = next(s for s in new.shapes if s.shape_type == 13)
    assert pic.image.blob == png_bytes("blue")  # embed re-pointed to a live relationship


def test_identical_images_from_two_slides_do_not_collide_in_target():
    src, slide = _source_with_picture_table_link()
    target = Presentation()
    layout = _pick_blank_layout(target)
    copy_slide_native(src, slide, target, layout)
    copy_slide_native(src, slide, target, layout)
    reopened = _roundtrip(target)
    blobs = [next(s for s in sl.shapes if s.shape_type == 13).image.blob for sl in reopened.slides]
    assert blobs[0] == blobs[1] == png_bytes("blue")


def test_slide_background_is_copied():
    src = Presentation()
    slide = src.slides.add_slide(src.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(1, 2, 3)
    target = Presentation()
    new = copy_slide_native(src, slide, target)
    assert str(new.background.fill.fore_color.rgb) == "010203"


def test_plain_table_is_supported_but_chart_is_not():
    src, slide = _source_with_picture_table_link()
    assert slide_has_unsupported_content(slide) is False

    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
    assert slide_has_unsupported_content(slide) is True
