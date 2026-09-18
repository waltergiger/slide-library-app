"""Builds the final .pptx from a list of chapters, each holding an ordered
list of slides picked from the library.

Native .pptx slides are copied in fully editable (pptx_copy.copy_slide_native);
PDF-sourced slides, and any .pptx slide with content the copier can't safely
handle (charts, SmartArt, embedded objects), come in as a full-bleed image of
that slide instead — the same trade-off shown in the mockup.
"""
from __future__ import annotations

import io
import logging

from pptx import Presentation
from pptx.util import Emu, Pt

from . import db, indexer
from .pptx_copy import copy_slide_native, slide_has_unsupported_content, _pick_blank_layout

log = logging.getLogger("slide-library.exporter")

SLIDE_WIDTH = Emu(12192000)   # 13.333 in — 16:9 widescreen
SLIDE_HEIGHT = Emu(6858000)   # 7.5 in
ACCENT = "2B3A55"


def build_deck(chapters: list[dict], add_dividers: bool = True) -> bytes:
    target = Presentation()
    target.slide_width = SLIDE_WIDTH
    target.slide_height = SLIDE_HEIGHT
    blank_layout = _pick_blank_layout(target)

    with indexer.PdfRenderCache() as cache:
        for chapter in chapters:
            name = (chapter.get("name") or "Untitled chapter").strip()
            slide_refs = chapter.get("slides", [])
            if add_dividers:
                _add_divider(target, blank_layout, name)
            for ref in slide_refs:
                _add_one_slide(target, blank_layout, ref, cache)

    buf = io.BytesIO()
    target.save(buf)
    return buf.getvalue()


def _add_one_slide(target: Presentation, blank_layout, ref: dict, cache) -> None:
    file_row = db.get_file(ref["file_id"])
    if file_row is None:
        return
    slide_index = ref["slide_index"]

    if file_row["ext"] == "pdf":
        _paste_image_slide(target, blank_layout, file_row, slide_index, cache)
        return

    try:
        source_prs = Presentation(file_row["path"])
        src_slide = source_prs.slides[slide_index]
        if slide_has_unsupported_content(src_slide):
            raise ValueError("slide has a chart/SmartArt/embedded object")
        copy_slide_native(source_prs, src_slide, target, blank_layout)
    except Exception as exc:  # noqa: BLE001
        log.info(
            "Falling back to image for %s slide %s: %s", file_row["path"], slide_index, exc
        )
        _paste_image_slide(target, blank_layout, file_row, slide_index, cache)


def _paste_image_slide(target: Presentation, blank_layout, file_row, slide_index: int, cache) -> None:
    slide = target.slides.add_slide(blank_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    png_bytes = indexer.render_full_size(
        file_row["path"], file_row["ext"], slide_index, target_width_px=1920, cache=cache
    )
    slide.shapes.add_picture(
        io.BytesIO(png_bytes), 0, 0, width=SLIDE_WIDTH, height=SLIDE_HEIGHT
    )


def _add_divider(target: Presentation, blank_layout, name: str) -> None:
    slide = target.slides.add_slide(blank_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)

    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)  # MSO_SHAPE.RECTANGLE = 1
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(ACCENT)
    bg.line.fill.background()
    bg.shadow.inherit = False

    box = slide.shapes.add_textbox(Emu(914400), Emu(2800000), SLIDE_WIDTH - Emu(1828800), Emu(1200000))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = name
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = _rgb("FFFFFF")


def _rgb(hex_str: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_str)
